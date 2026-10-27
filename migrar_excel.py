"""
Migra el archivo RUTAS_FRECUENTES_PESV.xlsx (formato de ficha por hoja)
hacia data/rutas.json, que es la base de datos que usa la app.

Uso:
    python migrar_excel.py RUTAS_FRECUENTES_PESV.xlsx
"""
import sys
import json
import re
from pathlib import Path
import openpyxl

CAMPOS_TEXTO = {
    "origen": "Origen",
    "destino": "Destino",
    "frecuencia": "Frecuencia",
    "recorrido": "Recorrido",
    "riesgos": "Riesgos",
    "controles": "Controles",
    "medidas_preventivas": "Medidas preventivas",
    "distancia": "Distancia aproximada",
    "tiempo": ["Tiempo promedio", "Tiempo estimado"],
    "punto_critico_general": "Punto critico",
}


def normaliza(txt):
    return str(txt).strip().lower() if txt is not None else ""


def migrar(path_excel):
    wb = openpyxl.load_workbook(path_excel, data_only=True)
    rutas = []

    for hoja in wb.sheetnames:
        ws = wb[hoja]
        datos = {clave: "" for clave in [
            "nombre", "cliente", "origen", "destino", "frecuencia",
            "recorrido", "riesgos", "controles", "medidas_preventivas",
            "distancia", "tiempo", "mapa_url"
        ]}
        datos["nombre"] = hoja
        puntos_criticos = []
        fila_tabla_riesgos = None

        filas = list(ws.iter_rows(values_only=False))
        for i, row in enumerate(filas):
            a = row[0].value if len(row) > 0 else None
            b = row[1].value if len(row) > 1 else None
            c = row[2].value if len(row) > 2 else None

            if a and normaliza(a).startswith("ruta"):
                datos["nombre"] = str(a).strip()
            if a and normaliza(a) == "nombre cliente":
                datos["cliente"] = str(b).strip() if b else ""

            for clave, etiquetas in CAMPOS_TEXTO.items():
                etiquetas_lista = etiquetas if isinstance(etiquetas, list) else [etiquetas]
                for etiqueta in etiquetas_lista:
                    if a and normaliza(a) == normaliza(etiqueta):
                        if clave == "punto_critico_general":
                            continue
                        datos[clave] = str(b).strip() if b else ""

            # Detecta el encabezado de la tabla de puntos críticos
            if a and normaliza(a) == "punto" and b and normaliza(b).startswith("riesgo"):
                fila_tabla_riesgos = i

            # Detecta el link de Google My Maps (única celda en col A, sin B/C)
            if a and isinstance(a, str) and "google.com/maps" in a:
                datos["mapa_url"] = a.strip()

        # Recorre filas debajo del encabezado de la tabla hasta encontrar vacío
        if fila_tabla_riesgos is not None:
            for row in filas[fila_tabla_riesgos + 1:]:
                a = row[0].value if len(row) > 0 else None
                b = row[1].value if len(row) > 1 else None
                c = row[2].value if len(row) > 2 else None
                if a is None and b is None and c is None:
                    break
                if a and "google.com/maps" in str(a):
                    break
                puntos_criticos.append({
                    "punto": str(a).strip() if a else "",
                    "riesgo": str(b).strip() if b else "",
                    "nivel": str(c).strip() if c else "Medio",
                })

        datos["puntos_criticos"] = puntos_criticos
        rutas.append(datos)

    return rutas


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python migrar_excel.py archivo.xlsx")
        sys.exit(1)

    path_excel = sys.argv[1]
    rutas = migrar(path_excel)

    out_dir = Path(__file__).parent / "data"
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / "rutas.json"

    if out_path.exists():
        respuesta = input(f"{out_path} ya existe. ¿Sobrescribir? (s/n): ")
        if respuesta.lower() != "s":
            print("Cancelado.")
            sys.exit(0)

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(rutas, f, ensure_ascii=False, indent=2)

    print(f"Migradas {len(rutas)} rutas -> {out_path}")
