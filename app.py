"""
App PESV - Rutas Frecuentes y Riesgos Viales
Gestión dinámica de fichas de ruta para el Plan Estratégico de Seguridad Vial.
"""
import json
import re
from pathlib import Path
from io import BytesIO

import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from utils_pesv import (
    convertir_a_embed, generar_pdf_ficha,
    cargar_consultas_estado, registrar_consulta,
    ruta_imagen_path, guardar_imagen_mapa, eliminar_imagen_mapa,
)
from config import ADMIN_PASSWORD
from ajustes import MODO_PUBLICO, AVISO_DATOS

DATA_PATH = Path(__file__).parent / "data" / "rutas.json"
NIVELES = ["Bajo", "Medio", "Alto"]
NIVEL_COLOR = {"Bajo": "🟢", "Medio": "🟡", "Alto": "🔴"}

_icono = Path(__file__).parent / "icon.png"
st.set_page_config(
    page_title="PESV - Rutas Frecuentes",
    page_icon=str(_icono) if _icono.exists() else "🚦",
    layout="wide",
)

# Ajustes para pantallas pequeñas (celular): menos márgenes, títulos más chicos, botones grandes
st.markdown("""
<style>
@media (max-width: 640px) {
    .block-container { padding: 1rem 0.8rem !important; }
    h1 { font-size: 1.5rem !important; }
    h2 { font-size: 1.2rem !important; }
    .stButton button, .stDownloadButton button,
    [data-testid="stFormSubmitButton"] button { width: 100%; min-height: 2.8rem; }
}
</style>
""", unsafe_allow_html=True)


# ---------------------------------------------------------------------------
# Persistencia
# ---------------------------------------------------------------------------
def cargar_rutas():
    if DATA_PATH.exists():
        with open(DATA_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    return []


def guardar_rutas(rutas):
    DATA_PATH.parent.mkdir(exist_ok=True)
    with open(DATA_PATH, "w", encoding="utf-8") as f:
        json.dump(rutas, f, ensure_ascii=False, indent=2)


if "rutas" not in st.session_state:
    st.session_state.rutas = cargar_rutas()


def ruta_vacia():
    return {
        "nombre": "", "cliente": "", "origen": "", "destino": "",
        "frecuencia": "", "recorrido": "", "riesgos": "", "controles": "",
        "medidas_preventivas": "", "distancia": "", "tiempo": "",
        "mapa_url": "", "mapa_imagen": "", "puntos_criticos": [],
    }


# ---------------------------------------------------------------------------
# Exportar a Excel (misma estructura del archivo original: 1 hoja por ruta)
# ---------------------------------------------------------------------------
def exportar_excel(rutas):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    titulo_font = Font(bold=True, size=11, color="FFFFFF")
    titulo_fill = PatternFill("solid", fgColor="1F4E78")
    campo_font = Font(bold=True)
    encabezado_font = Font(bold=True, color="FFFFFF")
    encabezado_fill = PatternFill("solid", fgColor="2E75B6")
    nivel_fill = {
        "Alto": PatternFill("solid", fgColor="F8CBAD"),
        "Medio": PatternFill("solid", fgColor="FFE699"),
        "Bajo": PatternFill("solid", fgColor="C6E0B4"),
    }
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    for ruta in rutas:
        nombre_hoja = re.sub(r'[\\/*?:\[\]]', '', ruta["nombre"])[:31] or "Ruta"
        # Evita nombres de hoja duplicados
        base, i = nombre_hoja, 1
        while nombre_hoja in wb.sheetnames:
            i += 1
            nombre_hoja = f"{base[:28]}_{i}"

        ws = wb.create_sheet(nombre_hoja)
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 55
        ws.column_dimensions["C"].width = 16

        ws.merge_cells("A1:C1")
        ws["A1"] = ruta["nombre"]
        ws["A1"].font = titulo_font
        ws["A1"].fill = titulo_fill
        ws["A1"].alignment = Alignment(horizontal="center")

        campos = [
            ("Cliente", ruta.get("cliente", "")),
            ("Origen", ruta.get("origen", "")),
            ("Destino", ruta.get("destino", "")),
            ("Frecuencia", ruta.get("frecuencia", "")),
            ("Recorrido", ruta.get("recorrido", "")),
            ("Riesgos generales", ruta.get("riesgos", "")),
            ("Controles", ruta.get("controles", "")),
            ("Medidas preventivas", ruta.get("medidas_preventivas", "")),
            ("Distancia aproximada", ruta.get("distancia", "")),
            ("Tiempo estimado", ruta.get("tiempo", "")),
        ]
        fila = 3
        for etiqueta, valor in campos:
            ws.cell(row=fila, column=1, value=etiqueta).font = campo_font
            c = ws.cell(row=fila, column=2, value=valor)
            c.alignment = wrap
            ws.cell(row=fila, column=1).border = border
            c.border = border
            fila += 1

        fila += 1
        for col, texto in zip("ABC", ["Punto", "Riesgo", "Nivel"]):
            c = ws[f"{col}{fila}"]
            c.value = texto
            c.font = encabezado_font
            c.fill = encabezado_fill
            c.border = border
        fila += 1
        for pc in ruta.get("puntos_criticos", []):
            ws.cell(row=fila, column=1, value=pc.get("punto", "")).border = border
            ws.cell(row=fila, column=2, value=pc.get("riesgo", "")).border = border
            nc = ws.cell(row=fila, column=3, value=pc.get("nivel", ""))
            nc.border = border
            nc.fill = nivel_fill.get(pc.get("nivel", ""), PatternFill())
            fila += 1

        if ruta.get("mapa_url"):
            fila += 1
            ws.cell(row=fila, column=1, value="Mapa (Google My Maps):").font = campo_font
            ws.cell(row=fila + 1, column=1, value=ruta["mapa_url"])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------
st.title("🚦 PESV — Rutas Frecuentes y Riesgos Viales")
st.subheader("Sistema de Gestión de Seguridad y Salud en el Trabajo")
"Identificación de rutas frecuentes y advertencia de riesgos a los conductores"

if "es_admin" not in st.session_state:
    st.session_state.es_admin = False

if MODO_PUBLICO:
    st.session_state.es_admin = False   # en la versión publicada no existe el modo administrador
else:
    with st.sidebar.expander("🔐 Acceso administrador", expanded=False):
        if st.session_state.es_admin:
            st.success("Sesión de administrador activa")
            if st.button("Cerrar sesión de administrador"):
                st.session_state.es_admin = False
                st.rerun()
        else:
            clave = st.text_input("Contraseña", type="password", key="clave_admin")
            if st.button("Ingresar"):
                if clave == ADMIN_PASSWORD:
                    st.session_state.es_admin = True
                    st.rerun()
                else:
                    st.error("Contraseña incorrecta.")

if st.session_state.es_admin:
    # Si algún botón pidió cambiar de pantalla (ej. "Editar"), se aplica aquí,
    # ANTES de dibujar el menú, para que el menú lateral quede sincronizado.
    if "_ir_a" in st.session_state:
        st.session_state["menu_admin"] = st.session_state.pop("_ir_a")
    menu = st.sidebar.radio(
        "Menú",
        [
            "🔎 Consultar ruta",
            "📋 Listado de rutas",
            "➕ Nueva ruta",
            "📝 Registro de consultas",
            "📊 Exportar / Respaldo",
        ],
        key="menu_admin",
    )
else:
    menu = "🔎 Consultar ruta"
    if not MODO_PUBLICO:
        st.sidebar.info("Modo conductor. Un administrador puede ingresar arriba para ver más opciones.")

rutas = st.session_state.rutas

# Si se sale de la pantalla de "Nueva ruta", se descarta cualquier edición pendiente
if menu != "➕ Nueva ruta":
    st.session_state.pop("editar_idx", None)

# Mensaje que debe sobrevivir a un st.rerun() (ej. "Ruta guardada")
if "_flash" in st.session_state:
    st.success(st.session_state.pop("_flash"))

# ---- CONSULTAR RUTA (vista del conductor) -----------------------------------
if menu == "🔎 Consultar ruta":
    st.subheader("🔎 Consultar ruta antes de salir")

    if not rutas:
        st.info("Todavía no hay rutas disponibles." if MODO_PUBLICO
                else "Todavía no hay rutas cargadas. Un administrador debe crear la primera en '➕ Nueva ruta'.")
    else:
        nombres_rutas = [r["nombre"] for r in rutas]

        with st.form("form_consulta"):
            c1, c2 = st.columns(2)
            conductor = c1.text_input("Nombre del conductor *")
            cedula = c2.text_input("Cédula *")
            ruta_sel_nombre = st.selectbox("Seleccione la ruta", nombres_rutas)
            consultar = st.form_submit_button("🔍 Consultar ruta", use_container_width=True)
        st.caption(AVISO_DATOS)

        if consultar:
            if not conductor.strip() or not cedula.strip():
                st.error("El nombre y la cédula son obligatorios para dejar registro de la consulta.")
            else:
                estado = registrar_consulta(conductor, ruta_sel_nombre, cedula)
                st.session_state["ruta_consultada"] = ruta_sel_nombre
                st.session_state["conductor_consulta"] = conductor
                if estado == "pendiente":
                    st.warning(
                        "Aquí tienes la información de la ruta, pero en este momento no se pudo "
                        "guardar el registro de tu consulta. Se reintentará automáticamente; "
                        "si el aviso se repite, informa a tu administrador."
                    )
                else:
                    st.success(f"Consulta registrada para {conductor} — {ruta_sel_nombre}")

        ruta_consultada_nombre = st.session_state.get("ruta_consultada")
        if ruta_consultada_nombre:
            ruta = next((r for r in rutas if r["nombre"] == ruta_consultada_nombre), None)
            if ruta:
                st.markdown("---")
                niveles_ruta = [pc.get("nivel", "") for pc in ruta.get("puntos_criticos", [])]
                nivel_max = "Alto" if "Alto" in niveles_ruta else ("Medio" if "Medio" in niveles_ruta else ("Bajo" if niveles_ruta else "—"))

                st.markdown(f"## {ruta['nombre']}")
                cA, cB, cC, cD = st.columns(4)
                cA.metric("Frecuencia", ruta.get("frecuencia", "—"))
                cB.metric("Distancia", ruta.get("distancia", "—"))
                cC.metric("Tiempo estimado", ruta.get("tiempo", "—"))
                cD.metric("Nivel de riesgo", f"{NIVEL_COLOR.get(nivel_max,'⚪')} {nivel_max}")

                st.markdown(f"**Origen:** {ruta.get('origen','—')}  \n**Destino:** {ruta.get('destino','—')}")
                st.markdown(f"**Recorrido:** {ruta.get('recorrido','—')}")
                if ruta.get("riesgos"):
                    st.markdown(f"**Riesgos generales:** {ruta.get('riesgos')}")

                if ruta.get("mapa_url"):
                    st.markdown("**Mapa de la ruta**")
                    embed_url = convertir_a_embed(ruta["mapa_url"])
                    st.components.v1.iframe(embed_url, height=480)

                if ruta.get("puntos_criticos"):
                    st.markdown("**Puntos críticos**")
                    for pc in ruta["puntos_criticos"]:
                        st.markdown(f"{NIVEL_COLOR.get(pc.get('nivel',''),'⚪')} **{pc.get('punto','')}** — {pc.get('riesgo','')}")

                if ruta.get("medidas_preventivas"):
                    st.markdown("**Medidas preventivas**")
                    for linea in ruta["medidas_preventivas"].split("\n"):
                        if linea.strip():
                            st.markdown(f"✔ {linea.strip()}")

                pdf_buffer = generar_pdf_ficha(ruta, st.session_state.get("conductor_consulta", ""))
                st.download_button(
                    "⬇️ Descargar ficha en PDF",
                    data=pdf_buffer,
                    file_name=f"ficha_{re.sub(r'[^A-Za-z0-9]+', '_', ruta['nombre'])}.pdf",
                    mime="application/pdf",
                )

# ---- LISTADO ---------------------------------------------------------------
elif menu == "📋 Listado de rutas":
    if not rutas:
        st.info("Todavía no hay rutas cargadas. Ve a '➕ Nueva ruta' para crear la primera.")
    else:
        col_f1, col_f2 = st.columns([2, 1])
        busqueda = col_f1.text_input("Buscar por nombre, origen o destino")
        niveles_filtro = col_f2.multiselect("Filtrar por nivel de riesgo presente", NIVELES)

        rutas_filtradas = []
        for r in rutas:
            texto = f"{r['nombre']} {r['origen']} {r['destino']}".lower()
            if busqueda and busqueda.lower() not in texto:
                continue
            if niveles_filtro:
                niveles_ruta = {pc.get("nivel") for pc in r.get("puntos_criticos", [])}
                if not niveles_ruta.intersection(niveles_filtro):
                    continue
            rutas_filtradas.append(r)

        st.write(f"**{len(rutas_filtradas)}** ruta(s) encontradas")

        for idx, ruta in enumerate(rutas_filtradas):
            i_real = rutas.index(ruta)
            niveles_ruta = [pc.get("nivel", "") for pc in ruta.get("puntos_criticos", [])]
            resumen_niveles = " ".join(NIVEL_COLOR.get(n, "") for n in sorted(set(niveles_ruta), key=lambda x: NIVELES.index(x) if x in NIVELES else 0))

            with st.expander(f"{ruta['nombre']}  {resumen_niveles}"):
                c1, c2 = st.columns(2)
                with c1:
                    st.markdown(f"**Cliente:** {ruta.get('cliente','—') or '—'}")
                    st.markdown(f"**Origen:** {ruta.get('origen','—')}")
                    st.markdown(f"**Destino:** {ruta.get('destino','—')}")
                    st.markdown(f"**Frecuencia:** {ruta.get('frecuencia','—')}")
                    st.markdown(f"**Distancia:** {ruta.get('distancia','—')}")
                    st.markdown(f"**Tiempo estimado:** {ruta.get('tiempo','—')}")
                with c2:
                    st.markdown(f"**Recorrido:** {ruta.get('recorrido','—')}")
                    st.markdown(f"**Riesgos generales:** {ruta.get('riesgos','—')}")
                    if ruta.get("controles"):
                        st.markdown(f"**Controles:** {ruta.get('controles')}")
                    if ruta.get("mapa_url"):
                        st.markdown(f"[🗺️ Ver mapa de la ruta]({ruta['mapa_url']})")
                    img_prev = ruta_imagen_path(ruta)
                    if img_prev:
                        st.image(str(img_prev), caption="Captura del mapa (va en el PDF)", width=300)
                    else:
                        st.caption("⚠️ Esta ruta no tiene captura del mapa: el PDF saldrá solo con el enlace.")

                if ruta.get("medidas_preventivas"):
                    st.markdown("**Medidas preventivas:**")
                    for linea in ruta["medidas_preventivas"].split("\n"):
                        if linea.strip():
                            st.markdown(f"- {linea.strip()}")

                if ruta.get("puntos_criticos"):
                    st.markdown("**Puntos críticos:**")
                    df_pc = pd.DataFrame(ruta["puntos_criticos"])
                    df_pc["Nivel"] = df_pc["nivel"].map(lambda n: f"{NIVEL_COLOR.get(n,'')} {n}")
                    st.dataframe(
                        df_pc[["punto", "riesgo", "Nivel"]].rename(
                            columns={"punto": "Punto", "riesgo": "Riesgo"}
                        ),
                        hide_index=True, use_container_width=True,
                    )

                b1, b2, b3 = st.columns(3)
                if b1.button("✏️ Editar", key=f"edit_{i_real}"):
                    st.session_state["editar_idx"] = i_real
                    st.session_state["_ir_a"] = "➕ Nueva ruta"
                    st.rerun()
                if b2.button("🗑️ Eliminar", key=f"del_{i_real}"):
                    eliminar_imagen_mapa(ruta.get("mapa_imagen", ""))
                    rutas.pop(i_real)
                    guardar_rutas(rutas)
                    st.rerun()
                pdf_buffer_admin = generar_pdf_ficha(ruta)
                b3.download_button(
                    "⬇️ PDF", data=pdf_buffer_admin,
                    file_name=f"ficha_{re.sub(r'[^A-Za-z0-9]+', '_', ruta['nombre'])}.pdf",
                    mime="application/pdf", key=f"pdf_{i_real}",
                )

# ---- NUEVA / EDITAR RUTA ----------------------------------------------------
elif menu == "➕ Nueva ruta":
    editar_idx = st.session_state.get("editar_idx")
    # Clave única por ruta: evita que los campos de una ruta se "mezclen" con otra
    clave_ruta = editar_idx if editar_idx is not None else "nueva"
    if editar_idx is not None:
        st.subheader(f"✏️ Editando: {rutas[editar_idx]['nombre']}")
        ruta_actual = rutas[editar_idx]
    else:
        st.subheader("➕ Nueva ruta")
        ruta_actual = ruta_vacia()

    # Va fuera del formulario para que, al cambiarlo, aparezcan o desaparezcan
    # al instante las filas de puntos críticos (dentro de un form no se actualiza).
    st.markdown("**Puntos críticos de la ruta**")
    n_puntos = st.number_input(
        "Número de puntos críticos (defínelo primero)", min_value=0, max_value=20,
        value=max(len(ruta_actual.get("puntos_criticos", [])), 1),
        key=f"npuntos_{clave_ruta}",
    )

    with st.form(f"form_ruta_{clave_ruta}", clear_on_submit=False):
        c1, c2 = st.columns(2)
        nombre = c1.text_input("Nombre de la ruta *", ruta_actual.get("nombre", ""))
        cliente = c2.text_input("Cliente / Sede", ruta_actual.get("cliente", ""))
        origen = c1.text_input("Origen", ruta_actual.get("origen", ""))
        destino = c2.text_input("Destino", ruta_actual.get("destino", ""))
        opciones_frec = ["Diaria", "Varias veces por semana", "Semanal", "Eventual", "Otra"]
        frec_actual = ruta_actual.get("frecuencia", "")
        if frec_actual and frec_actual not in opciones_frec:
            opciones_frec = [frec_actual] + opciones_frec  # conserva valores antiguos (ej. "Diaria (ejemplo)")
        frecuencia = c1.selectbox(
            "Frecuencia", opciones_frec,
            index=opciones_frec.index(frec_actual) if frec_actual in opciones_frec else 0,
        )
        distancia = c2.text_input("Distancia aproximada", ruta_actual.get("distancia", ""))
        tiempo = c1.text_input("Tiempo estimado", ruta_actual.get("tiempo", ""))
        mapa_url = c2.text_input("Link del mapa (Google My Maps)", ruta_actual.get("mapa_url", ""))

        st.markdown("**Captura del mapa para el PDF**")
        imagen_subida = st.file_uploader(
            "Sube una captura (PNG o JPG) del mapa de la ruta tal como se ve en Google My Maps",
            type=["png", "jpg", "jpeg"], key=f"img_{clave_ruta}",
        )
        img_actual = ruta_imagen_path(ruta_actual)
        quitar_imagen = False
        if img_actual:
            st.image(str(img_actual), caption="Captura actual (es la que sale en el PDF)", width=320)
            quitar_imagen = st.checkbox("Quitar la captura actual", key=f"quitar_{clave_ruta}")

        recorrido = st.text_area("Recorrido (trayecto paso a paso)", ruta_actual.get("recorrido", ""))
        riesgos = st.text_area("Riesgos generales de la vía", ruta_actual.get("riesgos", ""))
        controles = st.text_area("Controles existentes", ruta_actual.get("controles", ""))
        medidas = st.text_area(
            "Medidas preventivas (una por línea)",
            ruta_actual.get("medidas_preventivas", ""),
            height=120,
        )

        puntos_criticos = []
        for i in range(int(n_puntos)):
            pc_existente = (
                ruta_actual["puntos_criticos"][i]
                if i < len(ruta_actual.get("puntos_criticos", []))
                else {"punto": "", "riesgo": "", "nivel": "Medio"}
            )
            pc1, pc2, pc3 = st.columns([2, 3, 1])
            punto = pc1.text_input(f"Punto {i+1}", pc_existente["punto"], key=f"punto_{clave_ruta}_{i}")
            riesgo_pc = pc2.text_input(f"Riesgo {i+1}", pc_existente["riesgo"], key=f"riesgo_{clave_ruta}_{i}")
            nivel = pc3.selectbox(
                f"Nivel {i+1}", NIVELES,
                index=NIVELES.index(pc_existente["nivel"]) if pc_existente["nivel"] in NIVELES else 1,
                key=f"nivel_{clave_ruta}_{i}",
            )
            if punto or riesgo_pc:
                puntos_criticos.append({"punto": punto, "riesgo": riesgo_pc, "nivel": nivel})

        col_g1, col_g2 = st.columns(2)
        guardar = col_g1.form_submit_button("💾 Guardar ruta", use_container_width=True)
        cancelar = col_g2.form_submit_button("Cancelar", use_container_width=True)

    if cancelar:
        st.session_state.pop("editar_idx", None)
        st.session_state["_ir_a"] = "📋 Listado de rutas"
        st.rerun()

    if guardar:
        mapa_imagen = ruta_actual.get("mapa_imagen", "")
        error_imagen = None
        if imagen_subida is not None:
            try:
                nueva_img = guardar_imagen_mapa(imagen_subida.getvalue())
                eliminar_imagen_mapa(mapa_imagen)      # reemplaza la anterior
                mapa_imagen = nueva_img
            except Exception:
                error_imagen = "No se pudo leer el archivo como imagen. Sube un PNG o JPG válido."
        elif quitar_imagen and mapa_imagen:
            eliminar_imagen_mapa(mapa_imagen)
            mapa_imagen = ""

        if error_imagen:
            st.error(error_imagen)
        elif not nombre.strip():
            st.error("El nombre de la ruta es obligatorio.")
        else:
            nueva_ruta = {
                "nombre": nombre.strip(),
                "cliente": cliente.strip(),
                "origen": origen.strip(),
                "destino": destino.strip(),
                "frecuencia": frecuencia,
                "recorrido": recorrido.strip(),
                "riesgos": riesgos.strip(),
                "controles": controles.strip(),
                "medidas_preventivas": medidas.strip(),
                "distancia": distancia.strip(),
                "tiempo": tiempo.strip(),
                "mapa_url": mapa_url.strip(),
                "mapa_imagen": mapa_imagen,
                "puntos_criticos": puntos_criticos,
            }
            if editar_idx is not None:
                rutas[editar_idx] = nueva_ruta
            else:
                rutas.append(nueva_ruta)
            guardar_rutas(rutas)
            st.session_state.pop("editar_idx", None)
            st.session_state["_flash"] = f"Ruta «{nueva_ruta['nombre']}» guardada correctamente."
            st.session_state["_ir_a"] = "📋 Listado de rutas"
            st.rerun()

# ---- REGISTRO DE CONSULTAS (vista administrador) ------------------------------
elif menu == "📝 Registro de consultas":
    st.subheader("📝 Registro de consultas — evidencia PESV")
    st.caption("Registro de qué conductor consultó qué ruta y cuándo, antes de cada desplazamiento.")

    consultas, aviso_registro = cargar_consultas_estado()
    if aviso_registro:
        st.warning(aviso_registro)
    if not consultas:
        st.info("Todavía no se ha registrado ninguna consulta.")
    else:
        df = pd.DataFrame(consultas).drop(columns=["id"], errors="ignore")
        df = df.rename(columns={
            "conductor": "Conductor", "cedula": "Cédula",
            "ruta": "Ruta", "fecha_hora": "Fecha y hora",
        })
        df = df.sort_values("Fecha y hora", ascending=False)

        c1, c2 = st.columns(2)
        filtro_conductor = c1.text_input("Filtrar por conductor")
        filtro_ruta = c2.selectbox("Filtrar por ruta", ["Todas"] + sorted(df["Ruta"].unique().tolist()))

        df_filtrado = df.copy()
        if filtro_conductor:
            df_filtrado = df_filtrado[df_filtrado["Conductor"].str.contains(filtro_conductor, case=False, na=False)]
        if filtro_ruta != "Todas":
            df_filtrado = df_filtrado[df_filtrado["Ruta"] == filtro_ruta]

        st.write(f"**{len(df_filtrado)}** registro(s)")
        st.dataframe(df_filtrado, hide_index=True, use_container_width=True)

        csv = df_filtrado.to_csv(index=False).encode("utf-8-sig")
        st.download_button(
            "⬇️ Descargar registro (CSV)",
            data=csv,
            file_name="registro_consultas_pesv.csv",
            mime="text/csv",
        )

# ---- EXPORTAR ----------------------------------------------------------------
elif menu == "📊 Exportar / Respaldo":
    st.subheader("📊 Exportar / Respaldo")
    st.write(f"Actualmente hay **{len(rutas)}** ruta(s) registradas.")

    if rutas:
        excel_buffer = exportar_excel(rutas)
        st.download_button(
            "⬇️ Descargar Excel (una hoja por ruta)",
            data=excel_buffer,
            file_name="RUTAS_FRECUENTES_PESV.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        json_str = json.dumps(rutas, ensure_ascii=False, indent=2)
        st.download_button(
            "⬇️ Descargar respaldo JSON",
            data=json_str,
            file_name="rutas_respaldo.json",
            mime="application/json",
        )

        st.markdown("---")
        st.markdown("**Cargar respaldo JSON** (reemplaza todas las rutas actuales)")
        archivo = st.file_uploader("Selecciona un archivo .json", type=["json"])
        if archivo is not None:
            if st.button("Cargar y reemplazar"):
                nuevas = json.load(archivo)
                st.session_state.rutas = nuevas
                guardar_rutas(nuevas)
                st.success("Rutas cargadas correctamente.")
                st.rerun()
    else:
        st.info("No hay rutas para exportar todavía.")
